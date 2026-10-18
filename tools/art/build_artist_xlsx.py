"""Builds content/art/HANDOFF/Bhagavatam_paintings.xlsx from scenes.csv + PROMPT_GUIDE.md.
Run: py tools/art/build_artist_xlsx.py   (needs: pip install openpyxl)
"""
import csv
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "content", "art")
guide = open(os.path.join(ROOT, "PROMPT_GUIDE.md"), encoding="utf-8").read()
setup = next(b for b in re.findall(r"```\n(.*?)```", guide, re.S) if b.startswith("You paint chapter"))
style = setup.split("STYLE (always):")[1].split("CHARACTERS")[0].strip()
chars = setup.split("CHARACTERS (use exactly when named):")[1].strip()
rows = [r for r in csv.DictReader(open(os.path.join(ROOT, "scenes.csv"), encoding="utf-8")) if "DONE" not in r["title"]]

wb = Workbook()
ws = wb.active
ws.title = "Paintings"
head = ["File name to save as", "Chapter", "Prompt (copy this whole cell)", "Status", "Notes"]
ws.append(head)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="8A4B1F")
for r in rows:
    prompt = (f"Use the attached picture ch_0_1 only as the style reference, never copy its scene. Make ONE image, wide landscape 3:2.\n"
              f"SCENE: {r['scene']} TIME AND LIGHT: {r['light']}. FIGURES: {r['figures']}. MOOD: {r['mood']}.\n"
              f"{style}\nCHARACTERS: {chars}")
    ws.append([r["id"] + ".png", r["title"], prompt, "To do", ""])
for col, w in zip("ABCDE", (18, 34, 90, 12, 30)):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row[0].row].height = 60  # prompt cell is long; row stays compact, click the cell to copy
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
dv = DataValidation(type="list", formula1='"To do,Made,Approved,Redo"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"D2:D{ws.max_row}")

h = wb.create_sheet("How to use", 0)
for line in [
    "Bhagavatam chapter paintings: how to use this sheet",
    "",
    "1. Open the sheet 'Paintings'. Each row is one picture.",
    "2. Click the Prompt cell in column C, press Ctrl+C, paste into your image tool (ChatGPT, etc.).",
    "3. Attach the reference picture ch_0_1_REFERENCE.png every time (same folder as this file).",
    "4. If the picture is close, fix it with one short line in the same chat, e.g. 'Same scene. Move the figures into the right half.'",
    "5. Save the picture with the name in column A (exactly, for example ch_10_29.png).",
    "6. Set Status in column D: To do, Made, Approved or Redo. Use column E for notes.",
    "",
    "Check before saving: same soft gouache-on-paper look as the reference; the LEFT THIRD calm sky/water/mist; no writing, signature or border;",
    "hands and faces look right; no blood or gore; calm, kind faces.",
    "If a tool refuses a prompt: add 'a traditional devotional temple painting' once. Still refused? Write the file name in Notes and tell us.",
    "Start with the best-loved stories (Prahlada 7.8, Gajendra 8.2, Vamana 8.18-8.20, Putana 10.6, Kaliya 10.16, Govardhana 10.25, Rasa 10.29, Sudama 10.80) to settle the look.",
    "Skandha colours: 1 saffron, 2 golden yellow, 3 ochre, 4 leaf green, 5 jade, 6 teal, 7 vermilion, 8 deep blue, 9 amber, 10 indigo, 11 rose pink, 12 violet. Keep them.",
]:
    h.append([line])
h["A1"].font = Font(bold=True, size=14)
h.column_dimensions["A"].width = 150

out = os.path.join(ROOT, "HANDOFF", "Bhagavatam_paintings.xlsx")
wb.save(out)
print(len(rows), "rows ->", out)
