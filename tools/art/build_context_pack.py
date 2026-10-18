"""Builds content/art/HANDOFF/Bhagavatam_Art_Brief.docx and Bhagavatam_Art_Context.zip.
The .docx is the one file to give ChatGPT (or any image helper) so it understands the project and every scene.
Run: py tools/art/build_context_pack.py   (needs: pip install python-docx)
"""
import csv
import os
import re
import zipfile

from docx import Document
from docx.shared import Pt

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "content", "art")
H = os.path.join(ROOT, "HANDOFF")
guide = open(os.path.join(ROOT, "PROMPT_GUIDE.md"), encoding="utf-8").read()
setup = next(b for b in re.findall(r"```\n(.*?)```", guide, re.S) if b.startswith("You paint chapter"))
rows = [r for r in csv.DictReader(open(os.path.join(ROOT, "scenes.csv"), encoding="utf-8")) if "DONE" not in r["title"]]
rules = guide.split("### Rules for the scene line")[1].split("## Examples")[0]
fixing = guide.split("## Fixing a picture without starting again")[1]

d = Document()
d.styles["Normal"].font.name = "Calibri"
d.styles["Normal"].font.size = Pt(11)
d.add_heading("Bhagavatam chapter paintings: brief", 0)
d.add_paragraph(
    "This document is the context for making one painting per chapter of the Shrimad Bhagavat Mahapuran "
    "(Mahatmya and 12 Skandhas, 341 chapters). The pictures appear as the header of each chapter in an offline "
    "reading app. The app draws all lettering itself, so pictures must contain no writing. "
    "The attached picture ch_0_1 is the style reference: match its palette, brushwork, light, paper texture and detail, "
    "but never copy its scene. Chapter 0.1 is already done; 340 remain.")
d.add_heading("How to use this brief", 1)
for t in ["Read the Style, Format, Never and Characters sections once and keep them for every picture.",
          "When I send a chapter code (for example ch_10_29) or paste its prompt, make ONE image only, with no questions or suggestions.",
          "Use the scene table at the end to understand the story of that chapter. If a prompt and the table differ, the prompt wins."]:
    d.add_paragraph(t, style="List Bullet")
d.add_heading("Style, format, never, characters", 1)
for line in setup.split("\n"):
    if line.strip():
        d.add_paragraph(line.strip())
d.add_heading("Skandha colours", 1)
d.add_paragraph("Mahatmya antique gold; 1 saffron; 2 golden yellow; 3 ochre; 4 leaf green; 5 jade; 6 teal; 7 vermilion; "
                "8 deep blue; 9 amber; 10 indigo; 11 rose pink; 12 violet. Keep each Skandha's colour in the light and drapery.")
d.add_heading("Rules for every scene", 1)
for line in rules.split("\n"):
    if line.strip():
        d.add_paragraph(line.strip().replace("**", ""))
d.add_heading("Fixing a picture without starting again", 1)
for line in fixing.split("\n"):
    if line.strip() and not line.startswith("#"):
        d.add_paragraph(line.strip().lstrip("- "), style="List Bullet")
d.add_heading("All 340 scenes", 1)
d.add_paragraph("Columns: code, chapter title, scene, time and light, figures, mood.")
t = d.add_table(rows=1, cols=6)
t.style = "Light Grid Accent 1"
for c, x in zip(t.rows[0].cells, ["Code", "Chapter", "Scene", "Light", "Figures", "Mood"]):
    c.text = x
for r in rows:
    cells = t.add_row().cells
    for c, k in zip(cells, ["id", "title", "scene", "light", "figures", "mood"]):
        c.text = r[k]
        for p in c.paragraphs:
            for run in p.runs:
                run.font.size = Pt(8)
docx = os.path.join(H, "Bhagavatam_Art_Brief.docx")
d.save(docx)

zp = os.path.join(H, "Bhagavatam_Art_Context.zip")
with zipfile.ZipFile(zp, "w", zipfile.ZIP_DEFLATED) as z:
    z.write(docx, "Bhagavatam_Art_Brief.docx")
    z.write(os.path.join(H, "ch_0_1_REFERENCE.png"), "ch_0_1_REFERENCE.png")
    z.write(os.path.join(ROOT, "PROMPT_GUIDE.md"), "PROMPT_GUIDE.md")
    z.write(os.path.join(ROOT, "scenes.csv"), "scenes.csv")
    z.write(os.path.join(ROOT, "PROMPTS.md"), "PROMPTS_skandha_covers.md")
print(len(rows), "scenes;", os.path.getsize(docx) // 1024, "KB docx;", os.path.getsize(zp) // 1024, "KB zip")
